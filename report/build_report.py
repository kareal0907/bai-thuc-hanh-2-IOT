"""Sinh bao cao Word tu ket qua thuc nghiem.

Dau vao: report/results.json (gateway.evaluate), report/experiment.json (experiment_notes.py),
         data/preprocess_report.json (gateway.preprocess), report/figures/*.png
Chay:    python report/build_report.py   -> report/Bao_cao_Bai_2_IoT.docx
"""
import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
BASE = HERE.parent
FIG = HERE / "figures"
R = json.loads((HERE / "results.json").read_text(encoding="utf-8"))
X = json.loads((HERE / "experiment.json").read_text(encoding="utf-8"))
P = json.loads((BASE / "data" / "preprocess_report.json").read_text(encoding="utf-8"))
STUDENT = "Phạm Trung Huy"
STUDENT_ID = "B23DCAT132"
REPO_URL = "https://github.com/kareal0907/bai-thuc-hanh-2-IOT"
FONT = "Times New Roman"
ACCENT = RGBColor(0x1F, 0x3A, 0x68)

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.1
for level, size in ((1, 13.5), (2, 12)):
    st = doc.styles[f"Heading {level}"]
    st.font.name, st.font.size, st.font.bold = FONT, Pt(size), True
    st.font.color.rgb = ACCENT
    rfonts = st.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)  # font theme (Calibri Light) se de len font chi dinh
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    st.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.keep_with_next = True


def add_link(paragraph, url, text=None, size=None):
    """Chen hyperlink bam duoc (python-docx khong co san API cho hyperlink)."""
    r_id = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
                                         "hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rpr.append(el)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text or url
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return paragraph


def para(text="", bold_prefix=None, size=None, align=None, italic=False, after=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(1)
    return p


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def table(header, rows, widths, size=9.5, caption=None):
    if caption:
        c = para(caption, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, after=1)
        c.runs[0].bold = True
        c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold, run.font.size = True, Pt(size)
        shade(cell, "DCE6F2")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(size)
    for r_idx, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if r_idx == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                # bang ngan: giu tron tren mot trang
                p.paragraph_format.keep_with_next = len(rows) <= 8 and r_idx < len(t.rows) - 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, width_cm, caption, crop_left=0):
    if crop_left:  # bo sidebar cua anh dashboard cho gon
        from PIL import Image

        img = Image.open(path)
        path = HERE / "figures" / f"_crop_{Path(path).name}"
        img.crop((crop_left, 0, img.width, img.height)).save(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = para(caption, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, after=4)
    return c


def figures_side_by_side(items, width_cm, caption, crop_left=0):
    """Nhieu anh tren mot hang (bang khong vien), chung mot chu thich."""
    from PIL import Image

    t = doc.add_table(rows=1, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, (path, sub) in zip(t.rows[0].cells, items):
        img = Image.open(path)
        cropped = HERE / "figures" / f"_crop_{Path(path).name}"
        img.crop((crop_left, 0, img.width, img.height)).save(cropped)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(cropped), width=Cm(width_cm))
        s = cell.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = s.add_run(sub)
        r.italic, r.font.size = True, Pt(9)
    para(caption, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, after=4)


def ms(v):
    return "–" if v is None else f"{v:,.0f}".replace(",", ".")


def num(v, d=1):
    return f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------- Trang dau ----------------
t = para("BÁO CÁO BÀI THỰC HÀNH SỐ 2", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
t.runs[0].bold = True
t.runs[0].font.color.rgb = ACCENT
t = para("THU THẬP, LƯU TRỮ VÀ TIỀN XỬ LÝ DỮ LIỆU IoT", size=13.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
t.runs[0].bold = True
para(f"Học phần: IoT và Ứng dụng (INT14149)   ·   Sinh viên: {STUDENT}   ·   MSSV: {STUDENT_ID}",
     size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
add_link(para(f"Ngày thực nghiệm: {X['date']}   ·   Mã nguồn (GitHub): ", size=10.5,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=6), REPO_URL, size=10.5)

# ---------------- 1. Kien truc ----------------
doc.add_heading("1. Kiến trúc hệ thống", level=1)
para("Hệ thống gồm 4 lớp: thiết bị, truyền tải, gateway và lưu trữ/hiển thị (Hình 1). ESP32 mô phỏng trên Wokwi "
     "đọc DHT22 (nhiệt độ, độ ẩm), quang trở LDR (ánh sáng) và HC-SR04 (khoảng cách). Cứ 5 giây, thiết bị gửi một bản "
     "tin JSON gắn thời gian NTP và số thứ tự seq lên MQTT broker. Một thiết bị mô phỏng bằng Python (device_sim.py) "
     "dùng cùng định dạng và chủ động chèn lỗi để kiểm thử pipeline. Gateway Python (collector.py) subscribe, kiểm tra "
     "hợp lệ, loại trùng, phát hiện mất gói rồi ghi vào InfluxDB 2.7. Script preprocess.py tiền xử lý và ghi kết quả "
     "vào bucket riêng; dashboard Streamlit đọc cả ba bucket để giám sát.")
figure(FIG / "architecture.png", 15, "Hình 1. Sơ đồ kiến trúc pipeline thu thập – lưu trữ – tiền xử lý – giám sát")
table(["Thành phần", "Công nghệ", "Vai trò"], [
    ["Thiết bị", "ESP32 DevKit v1 (Wokwi), Arduino, PubSubClient, DHTesp",
     "Đo và publish JSON; NTP; bộ đệm store-and-forward 120 mẫu; LWT"],
    ["Broker", "EMQX public broker (MQTT 3.1.1, cổng 1883)", "Trung chuyển publish/subscribe, giữ phiên QoS1"],
    ["Gateway", "Python 3.14, paho-mqtt 2.1, influxdb-client 1.50", "Validate, dedup, đo độ trễ, spool khi DB lỗi"],
    ["Lưu trữ", "InfluxDB OSS 2.7.11 (time-series)", "3 bucket với retention khác nhau"],
    ["Tiền xử lý", "pandas 3.0, numpy, scikit-learn 1.9", "Làm sạch, outlier, resample, đặc trưng, chuẩn hóa"],
    ["Giám sát", "Streamlit 1.64 + Plotly", "Dashboard/app real-time, độ trễ, chất lượng, lưu trữ"],
], [2.4, 6.6, 8.0], caption="Bảng 1. Công nghệ sử dụng")

# ---------------- 2. Thu thap ----------------
doc.add_heading("2. Thu thập dữ liệu", level=1)
para("Topic được đặt theo cấu trúc phân cấp <root>/<site>/<device_id>/<loại>, ví dụ int14149/lab2/lab/esp32-01/telemetry. "
     "Nhờ vậy gateway có thể dùng wildcard int14149/lab2/+/+/telemetry để nhận dữ liệu từ nhiều thiết bị. Topic .../status "
     "mang trạng thái online/offline, gửi kèm cờ retained và làm Last Will. Payload mẫu:")
p = para('{"device_id":"esp32-01","ts":1790666221574,"seq":18,"temperature":25.00,"humidity":50.00,'
         '"light_lux":208.8,"distance_cm":100.86,"rssi":-81,"uptime_s":90}', size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT)
p.runs[0].font.name = "Consolas"
para("Gateway kiểm tra theo hai mức:")
bullet(" JSON hỏng, device_id không khớp topic, seq thiếu/không hợp lệ, hoặc không còn giá trị cảm biến nào hợp lệ "
       "→ loại cả bản tin.", "Mức cấu trúc:")
bullet(" giá trị null, sai kiểu hoặc ngoài dải đo vật lý (DHT22: −40…80 °C, 0…100 %RH; HC-SR04: 2…400 cm) → chỉ bỏ "
       "trường đó, vẫn giữ mẫu.", "Mức trường:")
para("Nếu ts bằng 0 hoặc lệch giờ quá ngưỡng (hơn 60 s ở tương lai, hoặc cũ hơn 15 phút), gateway dùng giờ của mình và "
     "gắn tag ts_source=gateway. Mọi sự cố đều được ghi vào iot_ops/data_quality để giám sát.")

# ---------------- 3. Schema ----------------
doc.add_heading("3. Thiết kế lưu trữ (InfluxDB)", level=1)
table(["Bucket (retention)", "Measurement", "Tags", "Fields"], [
    ["iot_raw (30 ngày)", "sensor_raw", "device_id, site, ts_source",
     "temperature, humidity, light_lux, distance_cm (float); seq, rssi, uptime_s (int)"],
    ["iot_processed (365 ngày)", "sensor_processed", "device_id, site, window, method",
     "<x>, <x>_rm, <x>_delta, <x>_z, <x>_mm; n_raw, n_imputed, n_outliers, completeness"],
    ["", "outlier_events", "device_id, field, method", "value"],
    ["iot_ops (7 ngày)", "pipeline_latency", "device_id, delivery", "ingest_ms, process_ms, write_ms, e2e_ms, seq"],
    ["", "data_quality", "device_id, kind", "count, seq"],
    ["", "device_status", "device_id", "status, online"],
], [3.3, 3.1, 3.9, 6.7], size=9, caption="Bảng 2. Schema lưu trữ")
para("Tag chỉ dùng cho thuộc tính ít giá trị, dùng để lọc (thiết bị, site). Giá trị đo và seq là field, vì đặt seq làm "
     "tag sẽ làm số series (cardinality) tăng không giới hạn. Timestamp là thời điểm đo tại thiết bị, độ chính xác ms. "
     "Retention được chia theo giá trị sử dụng của dữ liệu:")
bullet(" dữ liệu thô lớn, chỉ cần để truy vết ngắn hạn → 30 ngày;", "iot_raw:")
bullet(" dữ liệu 1 phút đã làm sạch, nhỏ hơn khoảng 12 lần → 1 năm;", "iot_processed:")
bullet(" số liệu vận hành (độ trễ, sự cố) → 7 ngày.", "iot_ops:")
para("Token của gateway chỉ có quyền đọc/ghi 3 bucket này (nguyên tắc quyền tối thiểu).")
table(["Tình huống", "Cơ chế xử lý"], [
    ["Trùng lặp (gửi lại, QoS1 redelivery)",
     "Cache LRU 5000 khóa (device_id, seq, ts). Nếu vẫn lọt, điểm cùng series + timestamp bị ghi đè, "
     "nên việc ghi là idempotent (lặp lại không tạo bản ghi mới)"],
    ["Mất gói / đến trễ / khởi động lại",
     "seq nhảy → gap (số gói mất). seq lùi và ts cũ hơn → out_of_order. seq lùi nhưng ts mới hơn → reset"],
    ["Thiết bị mất mạng", "Ring buffer 120 mẫu trên ESP32, gửi bù theo thứ tự (delivery=backfill)"],
    ["Gateway tạm dừng", "Client id cố định, clean_session=False, subscribe QoS1: broker giữ tin"],
    ["InfluxDB lỗi", "Ghi tạm vào spool.jsonl; kiểm tra lại DB mỗi 10 s và tự ghi bù"],
], [4.6, 12.4], size=9, caption="Bảng 3. Xử lý trùng lặp và mất mát dữ liệu")

# ---------------- 4. Tien xu ly ----------------
doc.add_heading("4. Tiền xử lý dữ liệu", level=1)
para("preprocess.py đọc sensor_raw theo khoảng thời gian (Flux + pivot) và xử lý riêng từng thiết bị theo các bước:")
bullet(" loại timestamp trùng; đưa dữ liệu về lưới đều 5 s để các mẫu thiếu hiện ra dưới dạng NaN.",
       "(1) Làm sạch:")
bullet(" Z-score trên cửa sổ trượt 60 mẫu (|z| > 3, mặc định) hoặc IQR (k = 1,5). Cả hai đều có ngưỡng phân tán tối "
       "thiểu (min_spread) để tránh chia cho 0 khi tín hiệu gần như không đổi. distance_cm không bị lọc, vì vật cản xuất "
       "hiện là sự kiện thật dạng bậc thang.", "(2) Outlier:")
bullet(" nội suy theo thời gian chỉ với khoảng trống ≤ 6 mẫu (30 s). Khoảng trống dài hơn giữ nguyên NaN để không tạo ra "
       "dữ liệu giả và được phản ánh qua chỉ số completeness.", "(3) Giá trị thiếu:")
bullet(" trung bình theo cửa sổ 1 phút; rolling mean 5 cửa sổ; delta giữa hai cửa sổ; chuẩn hóa Z-score "
       "(StandardScaler) và Min-Max (MinMaxScaler). Kết quả ghi vào iot_processed và xuất CSV.",
       "(4) Resample, đặc trưng, chuẩn hóa:")
rows = []
for dev, r in P.items():
    f = "temperature"
    rows.append([dev, r["rows_raw"], r["grid_slots"], r["missing_before"][f],
                 f"{sum(r['outliers_zscore'].values())} / {sum(r['outliers_iqr'].values())}",
                 r["imputed"][f], r["missing_after"][f], r["rows_processed"]])
table(["Thiết bị", "Mẫu thô", "Ô lưới 5 s", "Thiếu (nhiệt độ)", "Outlier Z / IQR", "Đã nội suy",
       "Còn thiếu", "Cửa sổ 1 phút"], rows, [2.3, 1.6, 1.9, 2.3, 2.5, 1.9, 1.7, 2.3], size=9,
      caption=f"Bảng 4. Kết quả tiền xử lý ({X['preprocess_window']})")
para(X["preprocess_comment"])
figure(FIG / "raw_vs_processed_sim-01.png", 14.5,
       "Hình 2. sim-01: dữ liệu thô (điểm xám), dữ liệu đã làm sạch trung bình 1 phút và rolling mean")

# ---------------- 5. Dashboard ----------------
doc.add_heading("5. Dashboard giám sát", level=1)
para("Dashboard Streamlit (python -m streamlit run gateway/dashboard.py) có 4 tab:")
bullet(" giá trị mới nhất, trạng thái online, tốc độ mẫu; tự làm mới 5 s; outlier đánh dấu ✕.", "Real-time:")
bullet(" so sánh dữ liệu thô, dữ liệu sạch và rolling mean; biểu đồ Z-score, delta, completeness.", "Dữ liệu đã xử lý:")
bullet(" p50/p95/p99, histogram, phân rã độ trễ theo chặng, thống kê sự cố.", "Độ trễ & chất lượng:")
bullet(" retention, số điểm, dung lượng; nút chạy tiền xử lý.", "Lưu trữ:")
figures_side_by_side(
    [(FIG / "dashboard_realtime_esp32-01.png", "(a) Real-time – esp32-01 (Wokwi)"),
     (FIG / "dashboard_latency_sim-01.png", "(b) Độ trễ & chất lượng – sim-01")],
    8.4, "Hình 3. Dashboard Streamlit. Ở (a), bước nhảy độ ẩm 50 → 71,5 % là lúc chỉnh thanh trượt DHT22 trên Wokwi",
    crop_left=340)

# ---------------- 6. Ket qua ----------------
doc.add_heading("6. Kết quả thực nghiệm", level=1)
para(X["setup_text"])
lat_rows = []
for label, key in X["latency_rows"]:
    d = X["latency"][key]
    lat_rows.append([label, d["n"], ms(d["ingest_p50"]), ms(d["write_mean"]), ms(d["e2e_p50"]),
                     ms(d["e2e_p95"]), ms(d["e2e_p99"]), ms(d["e2e_max"])])
table(["Nguồn (mẫu live)", "n", "Ingest p50", "Ghi DB TB", "E2E p50", "E2E p95", "E2E p99", "E2E max"],
      lat_rows, [4.6, 1.2, 1.9, 1.9, 1.9, 1.9, 1.9, 1.7], size=9, caption="Bảng 5. Độ trễ end-to-end (ms)")
figure(FIG / "latency.png", 14.5, "Hình 4. Phân bố và diễn biến độ trễ end-to-end")
q_rows = []
for dev, q in R["quality"].items():
    ev = q["events"]
    q_rows.append([dev, q["stored_samples"], q["expected_by_seq"] or "– (có reset)",
                   f"{num(q['delivery_ratio'] * 100, 1)} %" if q["delivery_ratio"] else "–",
                   ev.get("gap", 0), ev.get("duplicate", 0), ev.get("out_of_order", 0),
                   ev.get("reset", 0), sum(v for k, v in ev.items() if k in ("bad_json", "device_id_mismatch",
                                                                             "bad_seq", "no_valid_value")),
                   ev.get("null", 0) + ev.get("out_of_range", 0)])
table(["Thiết bị", "Mẫu lưu", "Kỳ vọng (seq)", "Tỉ lệ nhận", "Gap", "Trùng", "Trễ", "Reset", "Bị loại",
       "Trường lỗi"], q_rows, [2.1, 1.6, 2.2, 1.8, 1.2, 1.3, 1.2, 1.4, 1.6, 1.8], size=9,
      caption="Bảng 6. Chất lượng dữ liệu ghi nhận tại gateway")
table(["Kịch bản lỗi", "Kết quả quan sát"], X["fault_rows"], [4.5, 12.5], size=9,
      caption="Bảng 7. Kiểm thử khả năng chịu lỗi")
b = R["benchmark"]
st = R["storage"]
table(["Phép đo", "Kết quả"], [
    ["Ghi đồng bộ từng điểm (300 điểm)", f"{num(b['single_point_write']['points_per_s'], 0)} điểm/s "
                                         f"({num(b['single_point_write']['ms_per_write'], 2)} ms/lần ghi)"],
    ["Ghi theo lô 100 / 1.000 / 5.000 điểm (50.000 điểm)",
     " / ".join(num(b[f'batch_{k}_write']['points_per_s'], 0) for k in (100, 1000, 5000)) + " điểm/s"],
    ["Truy vấn 50.000 giá trị thô của 1 field", f"{num(b['query_raw_1_device_50k']['best_ms'], 0)} ms"],
    ["aggregateWindow 1 phút (5 field)", f"{num(b['query_mean_1m_window']['best_ms'], 0)} ms"],
    ["Pivot 5 field (50.000 dòng)", f"{num(b['query_pivot_5_fields']['best_ms'], 0)} ms"],
    ["RTT TCP tới broker (p50 / p95)", f"{num(R['broker_rtt_ms']['p50'], 0)} / {num(R['broker_rtt_ms']['p95'], 0)} ms"],
    ["Dung lượng iot_raw (sau khi snapshot sang TSM)",
     f"{num(st['iot_raw']['field_values'], 0)} giá trị field → TSM {num(st['iot_raw']['tsm_bytes'] / 1024, 1)} KB "
     f"(≈ {num(st['iot_raw']['tsm_bytes'] / st['iot_raw']['field_values'], 2)} B/giá trị sau nén); "
     f"file series/index cấp phát trước ≈ {num(st['iot_raw']['disk_bytes'] / 1024 / 1024, 0)} MB/bucket"],
    ["Benchmark 150.300 điểm × 5 field (bucket tạm)",
     f"TSM {num(b['bench_tsm_bytes'] / 1024, 0)} KB ≈ {num(b['bench_bytes_per_value'], 2)} B/giá trị"
     if b.get("bench_tsm_bytes") else "–"],
], [7.0, 10.0], size=9, caption="Bảng 8. Hiệu năng ghi/đọc InfluxDB (bucket tạm, máy tính cá nhân)")

# ---------------- 7. Phan tich ----------------
doc.add_heading("7. Phân tích độ trễ, lưu trữ và đề xuất cải tiến", level=1)
for head, text in X["analysis"]:
    para(text, bold_prefix=head + " ")
doc.add_heading("Đề xuất cải tiến", level=2)
for head, text in X["improvements"]:
    bullet(" " + text, head)

# ---------------- 8. Kho khan ----------------
doc.add_heading("8. Khó khăn và cách giải quyết", level=1)
table(["Khó khăn", "Cách giải quyết"], X["difficulties"], [6.0, 11.0], size=9)
doc.add_heading("9. Kết luận", level=1)
para(X["conclusion"].replace("nằm trong thư mục Bai2.", "được lưu tại repo GitHub công khai: "),
     align=WD_ALIGN_PARAGRAPH.LEFT)
add_link(doc.paragraphs[-1], REPO_URL)

out = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Bao_cao_Bai_2_IoT.docx"
doc.save(out)
print("saved", out)
